import html
import os
import shutil
from pathlib import Path

import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from .db import FILES, uid
from .drawing_standard import PRIMARY_STANDARD, PROFILE, drawing_standard_payload, validate_geometry

pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))


def valid_segments(geometry):
    result = []
    for item in geometry.get('segments', []):
        try:
            length = float(item['length_mm'])
            diameter = float(item['diameter_mm'])
        except (KeyError, TypeError, ValueError):
            return []
        if not (length > 0 and diameter > 0 and length < 100000 and diameter < 100000):
            return []
        result.append({'length_mm': length, 'diameter_mm': diameter,
                       'length_tolerance': str(item.get('length_tolerance') or ''),
                       'diameter_tolerance': str(item.get('diameter_tolerance') or ''),
                       'fit': str(item.get('fit') or ''),
                       'surface_roughness': str(item.get('surface_roughness') or ''),
                       'basis': item.get('basis', ''), 'note': item.get('note', '')})
    overall = geometry.get('overall_length_mm')
    if overall is not None:
        try:
            if abs(sum(s['length_mm'] for s in result) - float(overall)) > 0.05:
                return []
        except (TypeError, ValueError):
            return []
    return result


def profile_points(segments):
    x = 0.0
    top = [(0.0, segments[0]['diameter_mm'] / 2)]
    for i, segment in enumerate(segments):
        d = segment['diameter_mm'] / 2
        if top[-1][1] != d:
            top.append((x, d))
        x += segment['length_mm']
        top.append((x, d))
        if i + 1 < len(segments):
            top.append((x, segments[i + 1]['diameter_mm'] / 2))
    lower = [(px, -py) for px, py in reversed(top)]
    return top + lower + [top[0]]


from .drawing_iso import svg_preview_iso as svg_preview, draw_pdf_iso as draw_pdf, draw_dxf_iso as draw_dxf


from .dwg_converter import convert_dwg, ConversionError


from .native_rpc import native


@native("drawing")
def export_drawing(part):
    cad = (part.get('geometry') or {}).get('cad_document')
    if cad:
        paths = {key: cad.get(key, '') for key in ('drawing_pdf', 'drawing_dxf', 'drawing_dwg')}
        for value in paths.values():
            if value and (not Path(value).resolve().is_relative_to(FILES.resolve()) or not Path(value).is_file()):
                raise ValueError('CAD 版本文件无效，请重新上传')
        return paths
    folder = FILES / part['project_id'] / part['id'] / 'versions' / uid()
    folder.mkdir(parents=True, exist_ok=True)
    from .engineering_model import build_model
    from .cad_kernel import validate_solid
    import json
    model = build_model(part.get('geometry') or {}, part)
    model['cad_kernel'] = validate_solid(part.get('geometry') or {})
    (folder / 'engineering-model.json').write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding='utf-8')
    pdf = folder / 'drawing.pdf'
    dxf = folder / 'drawing.dxf'
    draw_pdf(part, pdf)
    complete = draw_dxf(part, dxf)
    dwg = None
    warning = ""
    if complete:
        try:
            dwg = convert_dwg(dxf)
        except Exception as exc:
            # PDF and DXF are the authoritative generated artifacts. LibreDWG
            # is an optional compatibility export and must not discard them.
            warning = f'DWG 附加导出失败（{type(exc).__name__}）：{exc}'
    return {'drawing_pdf': str(pdf), 'drawing_dxf': str(dxf) if complete else '', 'drawing_dwg': str(dwg) if dwg else '',
            'dwg_warning': warning, 'drawing_standard': drawing_standard_payload(), 'drawing_validation': validate_geometry(part)}


@native("process_pdf")
def export_process_pdf(project, part, process):
    process.setdefault('artifact_id', uid())
    folder = FILES / project['id'] / (part['id'] if part else 'assembly') / 'versions' / process['artifact_id']
    folder.mkdir(parents=True, exist_ok=True)
    output = folder / 'process.pdf'
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle
    from .manufacturing import manufacturing_summary
    style = ParagraphStyle('process-body', fontName='STSong-Light', fontSize=10, leading=16, wordWrap='CJK')
    heading = ParagraphStyle('process-heading', parent=style, fontSize=16, leading=24)
    def para(value, chosen=style):
        return Paragraph(html.escape(str(value or '')).replace('\n', '<br/>'), chosen)
    story = [para(process.get('title') or '工艺流程单', heading), Spacer(1,12),
             para(f"项目：{project['name']}　对象：{part['name'] if part else project['name']}　图号：{part['drawing_no'] if part else project['drawing_no']}"),
             Spacer(1,10)]
    if part:
        story += [para(line) for line in manufacturing_summary(part.get('geometry') or {})]
        story.append(Spacer(1,10))
    if process.get('summary'):
        story.extend([para(process['summary']), Spacer(1,10)])
    rows = [[para(x) for x in ('序号', '核心工序', '设备/工装', '工序要求', '检验要求')]]
    for step in process.get('steps') or []:
        rows.append([para(step.get(key,'')) for key in ('seq','operation','equipment','description','inspection')])
    table = Table(rows, colWidths=[40,130,140,415,385], repeatRows=1, splitByRow=1, splitInRow=1)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.grey),
                              ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF3F2')),
                              ('VALIGN',(0,0),(-1,-1),'TOP'),
                              ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story.append(table)
    for item in process.get('review_items') or []:
        story.extend([Spacer(1,6),para('待复核：'+str(item))])
    def footer(c, doc):
        c.setFont('STSong-Light',9)
        state='已审核' if process.get('approval_status')=='approved' else '工程草案 / 待审核'
        c.drawString(40,25,f'{state}　|　第 {doc.page} 页')
    SimpleDocTemplate(str(output), pagesize=landscape(A3), leftMargin=40, rightMargin=40,
                      topMargin=35, bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
    return output


@native("process_xlsx")
def export_process_xlsx(project, part, process):
    """An editable process card. Strings are forced to text to prevent formula injection."""
    process.setdefault('artifact_id', uid())
    folder = FILES / project['id'] / (part['id'] if part else 'assembly') / 'versions' / process['artifact_id']
    folder.mkdir(parents=True, exist_ok=True)
    output = folder / 'process.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '工艺流程单'
    dark = PatternFill('solid', fgColor='173E4B')
    pale = PatternFill('solid', fgColor='EAF3F2')
    line = Side(style='thin', color='C9D9DB')
    border = Border(left=line, right=line, top=line, bottom=line)

    def safe(value):
        if value is None:
            return ''
        value = str(value)
        return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value

    title = safe(process.get('title') or f"{part['name'] if part else project['name']}工艺流程单")
    sheet.merge_cells('A1:G1')
    sheet['A1'] = title
    sheet['A1'].font = Font(name='Microsoft YaHei', size=16, bold=True, color='FFFFFF')
    sheet['A1'].fill = dark
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 36
    metadata = [
        ('项目', project['name']), ('对象', part['name'] if part else project['name']),
        ('图号', part['drawing_no'] if part else project['drawing_no']),
        ('材料', process.get('material') or (part.get('material') if part else project['analysis'].get('material', ''))),
    ]
    for i, (label, value) in enumerate(metadata, 2):
        sheet.cell(i, 1, label).font = Font(bold=True, color='41626B')
        sheet.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
        sheet.cell(i, 2, safe(value))
        sheet.row_dimensions[i].height = 22
    headers = ['序号', '工序', '设备/工装', '工序说明', '检验要求', '依据', '复核']
    header_row = 7
    for col, label in enumerate(headers, 1):
        cell = sheet.cell(header_row, col, label)
        cell.fill = pale
        cell.font = Font(name='Microsoft YaHei', bold=True, color='244751')
        cell.alignment = Alignment(vertical='center', horizontal='center')
        cell.border = border
    sheet.row_dimensions[header_row].height = 28
    for index, step in enumerate(process.get('steps', []), 1):
        row_index = header_row + index
        values = [index, step.get('operation'), step.get('equipment'),
                  step.get('description'), step.get('inspection'),
                  step.get('basis'), '需复核' if step.get('review_required', True) else '已确认']
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col, index if col == 1 else safe(value))
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        longest = max(len(str(values[3] or '')), len(str(values[4] or '')))
        sheet.row_dimensions[row_index].height = min(108, max(42, 20 * ((longest + 36) // 37)))
    note_row = header_row + len(process.get('steps', [])) + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    sheet.cell(note_row, 1, '已审核 · 按已确认的图纸和工艺要求执行。' if process.get('approval_status') == 'approved' else '工程草案 · 图纸标注、推算和历史估算须由工程师核对后用于生产。')
    sheet.cell(note_row, 1).font = Font(color='A36438', italic=True)
    for offset, note in enumerate(process.get('review_items', []), 1):
        row_index = note_row + offset
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
        sheet.cell(row_index, 1, safe(f'待复核 {offset}：{note}'))
        sheet.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[row_index].height = min(60, max(28, len(str(note)) // 55 * 18 + 20))
    for col, width in enumerate([9, 20, 22, 60, 42, 22, 13], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = 'D8'
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:7'
    sheet.print_area = f'A1:G{max(note_row, note_row + len(process.get("review_items", [])))}'
    workbook.save(output)
    return output
